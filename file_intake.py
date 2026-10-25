#!/usr/bin/env python3
"""Read-only file intake with format adapters and traceable Markdown/JSON output."""

from __future__ import annotations

import argparse
import ast
import bz2
import csv
import gzip
import hashlib
import io
import json
import lzma
import mimetypes
import os
import re
import sqlite3
import sys
import tarfile
import zipfile
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path, PurePosixPath
from typing import Any, BinaryIO
from xml.etree import ElementTree


@dataclass(frozen=True)
class Limits:
    max_files: int = 5_000
    max_file_bytes: int = 50_000_000
    max_total_bytes: int = 500_000_000
    max_archive_members: int = 2_000
    max_archive_depth: int = 5
    max_database_rows: int = 100


class IntakeError(Exception):
    """An expected, user-actionable intake failure."""


MAX_PACKAGE_FILES = 5_000
MAX_PACKAGE_BYTES = 100_000_000
MAX_PACKAGE_PART_BYTES = 20_000_000


class TextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.hidden_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "noscript"}:
            self.hidden_depth += 1
        elif tag in {"p", "div", "br", "li", "tr", "h1", "h2", "h3"} and not self.hidden_depth:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript"} and self.hidden_depth:
            self.hidden_depth -= 1
        elif tag in {"p", "div", "li", "tr"} and not self.hidden_depth:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.hidden_depth and data.strip():
            self.parts.append(data.strip())


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _source_id(reference: str, digest: str) -> str:
    return "SRC-" + _sha256(f"{reference}\0{digest}".encode("utf-8"))[:20]


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _zip_xml(archive: zipfile.ZipFile, name: str) -> ElementTree.Element:
    try:
        info = archive.getinfo(name)
    except KeyError as exc:
        raise IntakeError(f"Required package part is missing: {name}") from exc
    if info.file_size > MAX_PACKAGE_PART_BYTES:
        raise IntakeError(f"Package XML part exceeds size limit: {name}")
    raw = archive.read(info)
    if b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
        raise IntakeError(f"Unsafe XML declarations are not accepted: {name}")
    return ElementTree.fromstring(raw)


def _read_limited(stream: BinaryIO, limit: int) -> bytes:
    data = stream.read(limit + 1)
    if len(data) > limit:
        raise IntakeError("FILE_SIZE_LIMIT_EXCEEDED")
    return data


class FileIntake:
    def __init__(self, output_dir: str | Path, limits: Limits | None = None) -> None:
        self.output_dir = Path(output_dir).expanduser().resolve()
        self.limits = limits or Limits()
        self.records: list[dict[str, Any]] = []
        self.total_bytes = 0
        self.intake_warnings: list[str] = []
        self.stopped = False
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.records_dir = self.output_dir / "readable"
        self.records_dir.mkdir(parents=True, exist_ok=True)

    def process(self, source: str | Path) -> dict[str, Any]:
        requested = Path(source).expanduser()
        if requested.is_symlink():
            raise IntakeError("SOURCE_SYMLINK_NOT_ALLOWED")
        path = requested.resolve(strict=True)
        if path.is_symlink():
            raise IntakeError("SOURCE_SYMLINK_NOT_ALLOWED")
        if path.is_file():
            self._process_path(path, path.name, path.as_posix(), None, None, 0)
        elif path.is_dir():
            for current, dirs, names in os.walk(path, followlinks=False):
                current_path = Path(current)
                dirs[:] = sorted(
                    name
                    for name in dirs
                    if not (current_path / name).is_symlink()
                    and not self._is_output_path(current_path / name)
                )
                for name in sorted(names):
                    if self.stopped:
                        break
                    item = current_path / name
                    if item.is_symlink() or not item.is_file() or self._is_output_path(item):
                        continue
                    self._process_path(
                        item,
                        item.name,
                        item.relative_to(path).as_posix(),
                        None,
                        None,
                        0,
                    )
                if self.stopped:
                    break
        else:
            raise IntakeError("SOURCE_MUST_BE_FILE_OR_DIRECTORY")
        return self._write_report(path.as_posix())

    def _is_output_path(self, path: Path) -> bool:
        try:
            path.resolve().relative_to(self.output_dir)
            return True
        except ValueError:
            return False

    def _check_budget(self, size: int) -> None:
        if len(self.records) > self.limits.max_files:
            raise IntakeError("MAX_FILES_EXCEEDED")
        if size > self.limits.max_file_bytes:
            raise IntakeError("FILE_SIZE_LIMIT_EXCEEDED")
        self.total_bytes += size
        if self.total_bytes > self.limits.max_total_bytes:
            raise IntakeError("MAX_TOTAL_BYTES_EXCEEDED")

    def _process_path(
        self,
        path: Path,
        name: str,
        reference: str,
        parent_id: str | None,
        member_path: str | None,
        depth: int,
    ) -> dict[str, Any]:
        if self.stopped:
            return {}
        if len(self.records) >= self.limits.max_files:
            self.stopped = True
            self.intake_warnings.append("MAX_FILES_EXCEEDED; remaining sources were skipped.")
            return {}
        try:
            size = path.stat().st_size
            if size > self.limits.max_file_bytes:
                return self._failure_record(
                    name, reference, parent_id, member_path, IntakeError("FILE_SIZE_LIMIT_EXCEEDED")
                )
            if self.total_bytes + size > self.limits.max_total_bytes:
                self.stopped = True
                self.intake_warnings.append("MAX_TOTAL_BYTES_EXCEEDED; remaining sources were skipped.")
                return self._failure_record(
                    name, reference, parent_id, member_path, IntakeError("MAX_TOTAL_BYTES_EXCEEDED")
                )
            with path.open("rb") as stream:
                data = _read_limited(stream, self.limits.max_file_bytes)
        except (OSError, IntakeError) as exc:
            return self._failure_record(name, reference, parent_id, member_path, exc)
        return self._process_bytes(data, name, reference, parent_id, member_path, depth, path)

    def _failure_record(
        self,
        name: str,
        reference: str,
        parent_id: str | None,
        member_path: str | None,
        error: Exception,
    ) -> dict[str, Any]:
        empty_hash = _sha256(b"")
        record = {
            "source_id": _source_id(reference, empty_hash),
            "content_sha256": None,
            "original_name": name,
            "original_extension": Path(name).suffix.lower(),
            "detected_format": "unknown",
            "source_path": reference,
            "size_bytes": None,
            "media_type": mimetypes.guess_type(name)[0],
            "captured_at": datetime.now(timezone.utc).isoformat(),
            "container_source_id": parent_id,
            "member_path": member_path,
            "status": "FAILED",
            "adapter_id": None,
            "adapter_version": None,
            "readable_output": None,
            "units": [],
            "evidence": [],
            "warnings": [],
            "errors": [f"{type(error).__name__}: {error}"],
            "metadata": {},
        }
        self.records.append(record)
        return record

    def _process_bytes(
        self,
        data: bytes,
        name: str,
        reference: str,
        parent_id: str | None,
        member_path: str | None,
        depth: int,
        source_path: Path | None = None,
    ) -> dict[str, Any]:
        digest = _sha256(data)
        record: dict[str, Any] = {
            "source_id": _source_id(reference, digest),
            "content_sha256": digest,
            "original_name": name,
            "original_extension": Path(name).suffix.lower(),
            "detected_format": "unknown",
            "source_path": reference,
            "size_bytes": len(data),
            "media_type": mimetypes.guess_type(name)[0],
            "captured_at": datetime.now(timezone.utc).isoformat(),
            "container_source_id": parent_id,
            "member_path": member_path,
            "status": "PROCESSING",
            "adapter_id": None,
            "adapter_version": "1",
            "readable_output": None,
            "units": [],
            "evidence": [{"kind": "SOURCE", "sha256": digest, "size_bytes": len(data)}],
            "warnings": [],
            "errors": [],
            "metadata": {},
        }
        self.records.append(record)
        try:
            self._check_budget(len(data))
            kind = self._detect(data, name)
            record["detected_format"] = kind
            record["media_type"] = self._media_type(kind, name)
            if kind == "zip":
                self._extract_zip(data, record, depth)
            elif kind == "tar":
                self._extract_tar(data, record, depth)
            elif kind in {"gzip", "bzip2", "xz"}:
                self._extract_compressed(data, record, kind, depth)
            else:
                adapter_id, readable, units, metadata, warnings = self._extract(
                    kind, data, name, source_path
                )
                record.update(
                    adapter_id=adapter_id,
                    readable_output=readable,
                    units=units,
                    metadata=metadata,
                    warnings=warnings,
                )
                record["status"] = "PARTIAL" if warnings else "EXTRACTED"
                if not readable and not units:
                    record["status"] = "UNSUPPORTED"
                    record["warnings"].append("No readable content was produced.")
            if record["readable_output"] and not record["units"]:
                record["units"] = [
                    {
                        "kind": "document",
                        "source_location": {"document": name},
                        "content": record["readable_output"],
                        "quality_status": "UNREVIEWED",
                    }
                ]
            for index, unit in enumerate(record["units"]):
                unit.setdefault("quality_status", "UNREVIEWED")
                unit.setdefault(
                    "unit_id",
                    "UNIT-" + _sha256((record["source_id"] + ":" + str(index)).encode())[:20],
                )
                unit.setdefault("evidence_sha256", _sha256(str(unit.get("content", "")).encode("utf-8")))
            self._write_readable(record)
        except (IntakeError, OSError, ValueError, UnicodeError, zipfile.BadZipFile, tarfile.TarError, sqlite3.Error, ElementTree.ParseError) as exc:
            record["status"] = "FAILED"
            record["errors"].append(f"{type(exc).__name__}: {exc}")
        return record

    @staticmethod
    def _detect(data: bytes, name: str) -> str:
        suffix = Path(name).suffix.lower()
        if data.startswith(b"%PDF-"):
            return "pdf"
        if data.startswith(b"SQLite format 3\0"):
            return "sqlite"
        if data.startswith(b"RIFF") and data[8:12] == b"WAVE":
            return "audio"
        if data.startswith(b"FORM") and data[8:12] in {b"AIFF", b"AIFC"}:
            return "audio"
        if data.startswith(b"RIFF") and data[8:12] == b"AVI ":
            return "video"
        if len(data) >= 12 and data[4:8] == b"ftyp":
            return "video"
        if data.startswith(b"\x1aE\xdf\xa3"):
            return "video"
        if data.startswith(b"ID3"):
            return "audio"
        if data.startswith(b"\x89PNG\r\n\x1a\n") or data.startswith(b"\xff\xd8\xff") or data.startswith((b"GIF87a", b"GIF89a", b"BM", b"II*\0", b"MM\0*")) or (len(data) > 12 and data[:4] == b"RIFF" and data[8:12] == b"WEBP"):
            return "image"
        if data.startswith(b"PK\x03\x04"):
            try:
                with zipfile.ZipFile(io.BytesIO(data)) as archive:
                    names = set(archive.namelist())
                if "word/document.xml" in names:
                    return "docx"
                if "xl/workbook.xml" in names:
                    return "xlsx"
                if "ppt/presentation.xml" in names:
                    return "pptx"
                if "mimetype" in names:
                    with zipfile.ZipFile(io.BytesIO(data)) as archive:
                        odf_type = archive.read("mimetype")[:128].decode("ascii", errors="ignore")
                    if odf_type == "application/vnd.oasis.opendocument.text":
                        return "odt"
                    if odf_type == "application/vnd.oasis.opendocument.spreadsheet":
                        return "ods"
                    if odf_type == "application/vnd.oasis.opendocument.presentation":
                        return "odp"
            except (zipfile.BadZipFile, OSError):
                pass
            return "zip"
        try:
            with tarfile.open(fileobj=io.BytesIO(data), mode="r:*"):
                return "tar"
        except (tarfile.TarError, OSError):
            pass
        known = {
            ".txt": "text",
            ".md": "markdown",
            ".rst": "text",
            ".log": "text",
            ".py": "python",
            ".pyw": "python",
            ".json": "json",
            ".jsonl": "jsonl",
            ".csv": "csv",
            ".tsv": "tsv",
            ".xml": "xml",
            ".html": "html",
            ".htm": "html",
            ".yaml": "yaml",
            ".yml": "yaml",
            ".toml": "toml",
            ".ini": "ini",
            ".cfg": "ini",
            ".sql": "sql",
            ".wav": "audio",
            ".mp3": "audio",
            ".m4a": "audio",
            ".flac": "audio",
            ".ogg": "audio",
            ".aiff": "audio",
            ".aif": "audio",
            ".mp4": "video",
            ".mov": "video",
            ".mkv": "video",
            ".avi": "video",
            ".webm": "video",
        }
        if suffix in known:
            return known[suffix]
        if data.startswith(b"\x1f\x8b"):
            return "gzip"
        if data.startswith(b"BZh"):
            return "bzip2"
        if data.startswith(b"\xfd7zXZ\x00"):
            return "xz"
        if data.startswith(b"7z\xbc\xaf\x27\x1c"):
            return "7z"
        if data.startswith(b"Rar!\x1a\x07"):
            return "rar"
        if b"\0" not in data[:4096]:
            try:
                data.decode("utf-8-sig")
                return "text"
            except UnicodeDecodeError:
                pass
        return "unknown"

    @staticmethod
    def _media_type(kind: str, name: str) -> str | None:
        if kind == "python":
            return "text/x-python"
        if kind == "sqlite":
            return "application/vnd.sqlite3"
        if kind in {"docx", "xlsx", "pptx"}:
            return {
                "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            }[kind]
        if kind == "audio":
            return mimetypes.guess_type(name)[0] or "audio/unknown"
        if kind == "video":
            return mimetypes.guess_type(name)[0] or "video/unknown"
        return mimetypes.guess_type(name)[0]

    def _extract(
        self, kind: str, data: bytes, name: str, source_path: Path | None
    ) -> tuple[str | None, str | None, list[dict[str, Any]], dict[str, Any], list[str]]:
        warnings: list[str] = []
        metadata: dict[str, Any] = {}
        units: list[dict[str, Any]] = []
        if kind in {"text", "markdown", "python", "yaml", "toml", "ini", "sql"}:
            text = data.decode("utf-8-sig", errors="replace")
            if "\ufffd" in text:
                warnings.append("Some bytes could not be decoded as UTF-8 and were replaced.")
            metadata["line_count"] = text.count("\n") + bool(text)
            if kind == "python":
                try:
                    tree = ast.parse(text, filename=name)
                    metadata["python_syntax"] = "valid"
                    metadata["imports"] = [
                        node.names[0].name
                        for node in ast.walk(tree)
                        if isinstance(node, ast.Import) and node.names
                    ] + [
                        node.module or ""
                        for node in ast.walk(tree)
                        if isinstance(node, ast.ImportFrom)
                    ]
                    metadata["symbols"] = [
                        {"kind": "class" if isinstance(node, ast.ClassDef) else "function", "name": node.name, "line": node.lineno}
                        for node in ast.walk(tree)
                        if isinstance(node, (ast.ClassDef, ast.FunctionDef, ast.AsyncFunctionDef))
                    ]
                except SyntaxError as exc:
                    metadata["python_syntax"] = "invalid"
                    warnings.append(f"Python syntax issue at line {exc.lineno}: {exc.msg}")
            return f"builtin.{kind}", text, units, metadata, warnings
        if kind in {"json", "jsonl"}:
            text = data.decode("utf-8-sig")
            if kind == "json":
                parsed = json.loads(text)
                metadata["top_level_type"] = type(parsed).__name__
                readable = json.dumps(parsed, ensure_ascii=False, indent=2)
            else:
                rows = [json.loads(line) for line in text.splitlines() if line.strip()]
                metadata["record_count"] = len(rows)
                readable = "\n".join(json.dumps(row, ensure_ascii=False, indent=2) for row in rows)
            return f"builtin.{kind}", readable, units, metadata, warnings
        if kind in {"csv", "tsv"}:
            delimiter = "\t" if kind == "tsv" else ","
            text = data.decode("utf-8-sig")
            rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
            metadata["row_count"] = max(0, len(rows) - 1)
            metadata["column_count"] = len(rows[0]) if rows else 0
            return f"builtin.{kind}", self._format_table(rows), units, metadata, warnings
        if kind == "xml":
            if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
                raise IntakeError("XML_DTD_AND_ENTITY_DECLARATIONS_NOT_ALLOWED")
            root = ElementTree.fromstring(data)
            metadata["root_element"] = _local(root.tag)
            return "builtin.xml", self._xml_text(root), units, metadata, warnings
        if kind == "html":
            parser = TextExtractor()
            parser.feed(data.decode("utf-8", errors="replace"))
            return "builtin.html", re.sub(r"\n{3,}", "\n\n", "\n".join(parser.parts)).strip(), units, metadata, warnings
        if kind == "docx":
            return self._extract_docx(data)
        if kind == "xlsx":
            return self._extract_xlsx(data)
        if kind == "pptx":
            return self._extract_pptx(data)
        if kind in {"odt", "ods", "odp"}:
            return self._extract_odf(data, kind)
        if kind == "pdf":
            return self._extract_pdf(data)
        if kind == "sqlite":
            if source_path is None:
                return None, None, [], {}, ["SQLite files inside archives require a dedicated in-memory database adapter."]
            return self._extract_sqlite(source_path)
        if kind == "image":
            return self._extract_image(data)
        if kind == "audio":
            return self._extract_audio(data, name)
        if kind == "video":
            return self._extract_video(data, name)
        if kind in {"7z", "rar", "gzip"}:
            return None, None, [], {"container_format": kind}, [f"No safe {kind} adapter is installed."]
        if kind in {"bzip2", "xz"}:
            return None, None, [], {"container_format": kind}, [f"No safe {kind} adapter is installed."]
        return None, None, [], {"extension": Path(name).suffix.lower()}, ["No adapter is registered for this format."]

    @staticmethod
    def _format_table(rows: list[list[str]]) -> str:
        if not rows:
            return "(empty table)"
        width = max(len(row) for row in rows)
        normalized = [row + [""] * (width - len(row)) for row in rows]
        cells = [[cell.replace("|", "\\|").replace("\n", " ") for cell in row] for row in normalized]
        header = cells[0]
        separator = ["---"] * width
        body = cells[1:] if len(cells) > 1 else []
        return "\n".join(["| " + " | ".join(header) + " |", "| " + " | ".join(separator) + " |"] + ["| " + " | ".join(row) + " |" for row in body])

    @staticmethod
    def _xml_text(root: ElementTree.Element) -> str:
        lines: list[str] = []
        for node in root.iter():
            if node.text and node.text.strip():
                lines.append(node.text.strip())
        return "\n".join(lines)

    def _extract_docx(self, data: bytes) -> tuple[str, str, list[dict[str, Any]], dict[str, Any], list[str]]:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            self._validate_office_package(archive)
            root = _zip_xml(archive, "word/document.xml")
            paragraphs = []
            for paragraph in root.iter():
                if _local(paragraph.tag) != "p":
                    continue
                text = "".join(node.text or "" for node in paragraph.iter() if _local(node.tag) == "t")
                if text.strip():
                    paragraphs.append(text)
            media = [name for name in archive.namelist() if name.startswith("word/media/")]
        readable = "\n\n".join(paragraphs)
        warnings = [f"{len(media)} embedded media item(s) are referenced but not OCR-processed."] if media else []
        return "builtin.docx", readable, [], {"paragraph_count": len(paragraphs), "media_members": media}, warnings

    def _extract_xlsx(self, data: bytes) -> tuple[str, str, list[dict[str, Any]], dict[str, Any], list[str]]:
        try:
            from openpyxl import load_workbook
        except ImportError:
            return None, None, [], {}, ["Install openpyxl to read XLSX workbooks."]
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            self._validate_office_package(archive)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        sections: list[str] = []
        sheet_names = workbook.sheetnames
        for sheet in workbook.worksheets:
            sections.append(f"## Sheet: {sheet.title}")
            rows = []
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                while values and not values[-1]:
                    values.pop()
                if values:
                    rows.append(values)
            sections.append(self._format_table(rows))
        workbook.close()
        return "openpyxl.xlsx", "\n\n".join(sections), [], {"sheet_names": sheet_names}, []

    def _extract_pptx(self, data: bytes) -> tuple[str, str, list[dict[str, Any]], dict[str, Any], list[str]]:
        slides: list[tuple[int, str]] = []
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            self._validate_office_package(archive)
            names = [name for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)]
            for name in names:
                match = re.search(r"slide(\d+)\.xml", name)
                root = _zip_xml(archive, name)
                paragraphs = []
                for paragraph in root.iter():
                    if _local(paragraph.tag) == "p":
                        text = "".join(node.text or "" for node in paragraph.iter() if _local(node.tag) == "t")
                        if text.strip():
                            paragraphs.append(text)
                slides.append((int(match.group(1)) if match else 0, "\n".join(paragraphs)))
        slides.sort()
        readable = "\n\n".join(f"## Slide {number}\n{text}" for number, text in slides)
        units = [{"kind": "slide", "source_location": {"slide": number}, "content": text} for number, text in slides]
        return "builtin.pptx", readable, units, {"slide_count": len(slides)}, []

    def _extract_odf(self, data: bytes, kind: str) -> tuple[str, str, list[dict[str, Any]], dict[str, Any], list[str]]:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            self._validate_office_package(archive)
            root = _zip_xml(archive, "content.xml")
            media = [name for name in archive.namelist() if name.startswith("Pictures/")]
        paragraphs = []
        for node in root.iter():
            if _local(node.tag) in {"p", "h"}:
                text = "".join(child.text or "" for child in node.iter() if child is not node)
                if text.strip():
                    paragraphs.append(text.strip())
        warnings = [f"{len(media)} embedded media item(s) are referenced but not OCR-processed."] if media else []
        return f"builtin.{kind}", "\n\n".join(paragraphs), [], {"paragraph_count": len(paragraphs), "media_members": media}, warnings

    @staticmethod
    def _validate_office_package(archive: zipfile.ZipFile) -> None:
        members = archive.infolist()
        if len(members) > MAX_PACKAGE_FILES:
            raise IntakeError("OFFICE_PACKAGE_MEMBER_LIMIT_EXCEEDED")
        if sum(member.file_size for member in members) > MAX_PACKAGE_BYTES:
            raise IntakeError("OFFICE_PACKAGE_EXPANDED_SIZE_LIMIT_EXCEEDED")

    def _extract_pdf(self, data: bytes) -> tuple[str, str, list[dict[str, Any]], dict[str, Any], list[str]]:
        try:
            from pypdf import PdfReader
        except ImportError:
            return None, None, [], {}, ["Install pypdf to extract embedded PDF text; OCR is a separate optional capability."]
        reader = PdfReader(io.BytesIO(data), strict=False)
        if reader.is_encrypted:
            raise IntakeError("PDF_IS_ENCRYPTED")
        pages: list[str] = []
        units: list[dict[str, Any]] = []
        warnings: list[str] = []
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            pages.append(f"## Page {number}\n{text}")
            units.append({"kind": "page", "source_location": {"page": number}, "content": text})
            if not text.strip():
                warnings.append(f"Page {number} has no embedded text; OCR was not performed.")
        return "pypdf", "\n\n".join(pages), units, {"page_count": len(reader.pages)}, warnings

    def _extract_image(self, data: bytes) -> tuple[str, str, list[dict[str, Any]], dict[str, Any], list[str]]:
        try:
            from PIL import Image
        except ImportError:
            return None, None, [], {}, ["Install Pillow to read image metadata."]
        image = Image.open(io.BytesIO(data))
        metadata = {"format": image.format, "width": image.width, "height": image.height, "mode": image.mode}
        warnings = []
        text = ""
        try:
            import pytesseract

            text = pytesseract.image_to_string(image)
            if not text.strip():
                warnings.append("OCR ran but did not detect readable text.")
        except ImportError:
            warnings.append("OCR is not configured; image pixels were not converted to text.")
        except (OSError, RuntimeError, ValueError) as exc:
            warnings.append(f"OCR failed: {type(exc).__name__}: {exc}")
        readable = (
            f"Image metadata\n"
            f"- Format: {image.format}\n"
            f"- Dimensions: {image.width} x {image.height}\n"
            f"- Mode: {image.mode}"
        )
        if text.strip():
            readable += "\n\n## OCR text\n\n" + text.strip()
        elif warnings:
            readable += "\n- OCR: unavailable"
        return "pillow.image_metadata", readable, [], metadata, warnings

    @staticmethod
    def _extract_audio(data: bytes, name: str) -> tuple[str, str, list[dict[str, Any]], dict[str, Any], list[str]]:
        metadata: dict[str, Any] = {"container": Path(name).suffix.lower().lstrip(".") or "unknown"}
        if data.startswith(b"RIFF") and data[8:12] == b"WAVE":
            import wave

            with wave.open(io.BytesIO(data), "rb") as audio:
                frames = audio.getnframes()
                rate = audio.getframerate()
                metadata.update(
                    channels=audio.getnchannels(),
                    sample_width_bytes=audio.getsampwidth(),
                    sample_rate_hz=rate,
                    duration_seconds=round(frames / rate, 3) if rate else None,
                )
        else:
            metadata["size_bytes"] = len(data)
        readable = "Audio metadata\n" + "\n".join(f"- {key}: {value}" for key, value in metadata.items())
        return "builtin.audio_metadata", readable, [], metadata, ["Speech transcription is not configured."]

    @staticmethod
    def _extract_video(data: bytes, name: str) -> tuple[str, str, list[dict[str, Any]], dict[str, Any], list[str]]:
        container = "mp4" if len(data) >= 12 and data[4:8] == b"ftyp" else Path(name).suffix.lower().lstrip(".") or "unknown"
        metadata = {"container": container, "size_bytes": len(data)}
        readable = f"Video metadata\n- Container: {container}\n- Size: {len(data)} bytes"
        return "builtin.video_container", readable, [], metadata, ["Video frames and audio are not decoded; install and configure media adapters."]

    def _extract_sqlite(self, path: Path) -> tuple[str, str, list[dict[str, Any]], dict[str, Any], list[str]]:
        uri = path.as_uri() + "?mode=ro"
        connection = sqlite3.connect(uri, uri=True)
        try:
            connection.execute("PRAGMA query_only = ON")
            tables = connection.execute(
                "SELECT name, sql FROM sqlite_master WHERE type = 'table' AND name NOT LIKE 'sqlite_%' ORDER BY name"
            ).fetchall()
            sections: list[str] = []
            metadata_tables = []
            for table_name, create_sql in tables:
                escaped = table_name.replace('"', '""')
                columns = connection.execute(f'PRAGMA table_info("{escaped}")').fetchall()
                count = connection.execute(f'SELECT COUNT(*) FROM "{escaped}"').fetchone()[0]
                cursor = connection.execute(f'SELECT * FROM "{escaped}" LIMIT ?', (self.limits.max_database_rows,))
                rows = [list(item) for item in cursor.fetchall()]
                header = [column[1] for column in columns]
                sections.extend([f"## Table: {table_name}", f"Schema: {create_sql or '(unavailable)'}", f"Rows: {count}", self._format_table([header] + [[str(value) if value is not None else "" for value in row] for row in rows])])
                metadata_tables.append({"name": table_name, "row_count": count, "columns": header, "preview_rows": len(rows)})
            return "sqlite.readonly", "\n\n".join(sections), [], {"tables": metadata_tables}, []
        finally:
            connection.close()

    def _extract_zip(self, data: bytes, record: dict[str, Any], depth: int) -> None:
        if depth >= self.limits.max_archive_depth:
            raise IntakeError("MAX_ARCHIVE_DEPTH_EXCEEDED")
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > self.limits.max_archive_members:
                raise IntakeError("MAX_ARCHIVE_MEMBERS_EXCEEDED")
            record["adapter_id"] = "builtin.zip"
            record["metadata"] = {"member_count": len(members)}
            record["readable_output"] = "\n".join(
                f"{'DIR' if info.is_dir() else 'FILE'} | {info.file_size} bytes | {info.filename}"
                for info in members
            )
            children = []
            for info in members:
                if info.is_dir():
                    continue
                if len(self.records) >= self.limits.max_files:
                    record["warnings"].append("MAX_FILES_EXCEEDED; remaining archive members were skipped.")
                    self.stopped = True
                    break
                pure = PurePosixPath(info.filename)
                if pure.is_absolute() or ".." in pure.parts or "\\" in info.filename:
                    record["warnings"].append(f"Unsafe archive path skipped: {info.filename}")
                    continue
                if info.file_size > self.limits.max_file_bytes:
                    children.append(self._failure_record(info.filename, f"{record['source_path']}!/{info.filename}", record["source_id"], info.filename, IntakeError("FILE_SIZE_LIMIT_EXCEEDED")))
                    continue
                if self.total_bytes + info.file_size > self.limits.max_total_bytes:
                    record["warnings"].append("MAX_TOTAL_BYTES_EXCEEDED; remaining archive members were skipped.")
                    break
                if info.flag_bits & 0x1:
                    children.append(self._failure_record(info.filename, f"{record['source_path']}!/{info.filename}", record["source_id"], info.filename, IntakeError("ARCHIVE_MEMBER_ENCRYPTED")))
                    continue
                try:
                    member_data = archive.read(info)
                except (OSError, RuntimeError, zipfile.BadZipFile) as exc:
                    children.append(self._failure_record(info.filename, f"{record['source_path']}!/{info.filename}", record["source_id"], info.filename, exc))
                    continue
                child = self._process_bytes(
                    member_data,
                    pure.name,
                    f"{record['source_path']}!/{info.filename}",
                    record["source_id"],
                    info.filename,
                    depth + 1,
                )
                children.append(child)
            record["metadata"]["processed_members"] = len(children)
            if any(child["status"] not in {"EXTRACTED", "PARTIAL"} for child in children):
                record["warnings"].append("One or more archive members were not fully extracted.")
            if record["warnings"]:
                record["status"] = "PARTIAL"
            else:
                record["status"] = "EXTRACTED"

    def _extract_tar(self, data: bytes, record: dict[str, Any], depth: int) -> None:
        if depth >= self.limits.max_archive_depth:
            raise IntakeError("MAX_ARCHIVE_DEPTH_EXCEEDED")
        with tarfile.open(fileobj=io.BytesIO(data), mode="r:*") as archive:
            members = archive.getmembers()
            if len(members) > self.limits.max_archive_members:
                raise IntakeError("MAX_ARCHIVE_MEMBERS_EXCEEDED")
            record["adapter_id"] = "builtin.tar"
            record["metadata"] = {"member_count": len(members)}
            record["readable_output"] = "\n".join(
                f"{'DIR' if member.isdir() else 'FILE' if member.isfile() else 'SKIPPED'} | {member.size} bytes | {member.name}"
                for member in members
            )
            children = []
            for member in members:
                if len(self.records) >= self.limits.max_files:
                    record["warnings"].append("MAX_FILES_EXCEEDED; remaining archive members were skipped.")
                    self.stopped = True
                    break
                pure = PurePosixPath(member.name)
                if not member.isfile():
                    if not member.isdir():
                        record["warnings"].append(f"Non-regular tar member skipped: {member.name}")
                    continue
                if pure.is_absolute() or ".." in pure.parts or "\\" in member.name:
                    record["warnings"].append(f"Unsafe archive path skipped: {member.name}")
                    continue
                if member.size > self.limits.max_file_bytes:
                    children.append(self._failure_record(member.name, f"{record['source_path']}!/{member.name}", record["source_id"], member.name, IntakeError("FILE_SIZE_LIMIT_EXCEEDED")))
                    continue
                if self.total_bytes + member.size > self.limits.max_total_bytes:
                    record["warnings"].append("MAX_TOTAL_BYTES_EXCEEDED; remaining archive members were skipped.")
                    break
                stream = archive.extractfile(member)
                if stream is None:
                    children.append(self._failure_record(member.name, f"{record['source_path']}!/{member.name}", record["source_id"], member.name, IntakeError("TAR_MEMBER_UNREADABLE")))
                    continue
                with stream:
                    member_data = _read_limited(stream, self.limits.max_file_bytes)
                children.append(
                    self._process_bytes(
                        member_data,
                        pure.name,
                        f"{record['source_path']}!/{member.name}",
                        record["source_id"],
                        member.name,
                        depth + 1,
                    )
                )
            record["metadata"]["processed_members"] = len(children)
            record["status"] = "PARTIAL" if record["warnings"] or any(child["status"] not in {"EXTRACTED", "PARTIAL"} for child in children) else "EXTRACTED"
            if any(child["status"] not in {"EXTRACTED", "PARTIAL"} for child in children):
                record["warnings"].append("One or more archive members were not fully extracted.")

    def _extract_compressed(
        self, data: bytes, record: dict[str, Any], kind: str, depth: int
    ) -> None:
        if depth >= self.limits.max_archive_depth:
            raise IntakeError("MAX_ARCHIVE_DEPTH_EXCEEDED")
        if kind == "gzip":
            stream: BinaryIO = gzip.GzipFile(fileobj=io.BytesIO(data))
        elif kind == "bzip2":
            stream = bz2.BZ2File(io.BytesIO(data))
        else:
            stream = lzma.LZMAFile(io.BytesIO(data))
        with stream:
            content = _read_limited(stream, self.limits.max_file_bytes)
        original_name = record["original_name"]
        member_name = re.sub(r"(?i)\.(?:gz|bz2|xz)$", "", original_name) or f"{original_name}.content"
        child = self._process_bytes(
            content,
            member_name,
            f"{record['source_path']}!/{member_name}",
            record["source_id"],
            member_name,
            depth + 1,
        )
        record["adapter_id"] = f"builtin.{kind}"
        record["metadata"] = {"decompressed_member": member_name}
        record["readable_output"] = f"Decompressed member: {member_name}"
        record["metadata"]["child_source_id"] = child["source_id"]
        if child["status"] in {"EXTRACTED", "PARTIAL"}:
            record["status"] = "EXTRACTED" if child["status"] == "EXTRACTED" else "PARTIAL"
        else:
            record["status"] = "PARTIAL"
            record["warnings"].append(f"Decompressed content status: {child['status']}")

    def _write_readable(self, record: dict[str, Any]) -> None:
        if not record["readable_output"] and not record["units"]:
            return
        target = self.records_dir / f"{record['source_id']}.md"
        if target.exists():
            return
        body = record["readable_output"] or ""
        target.write_text(
            f"# {record['original_name']}\n\n"
            f"- Source ID: `{record['source_id']}`\n"
            f"- Format: `{record['detected_format']}`\n"
            f"- SHA-256: `{record['content_sha256']}`\n"
            f"- Status: `{record['status']}`\n\n---\n\n{body}\n",
            encoding="utf-8",
        )
        record["readable_output_path"] = target.relative_to(self.output_dir).as_posix()

    def _write_report(self, source: str) -> dict[str, Any]:
        report = {
            "schema_version": "1.0",
            "source": source,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "limits": asdict(self.limits),
            "warnings": self.intake_warnings,
            "summary": {
                "sources": len(self.records),
                "extracted": sum(row["status"] == "EXTRACTED" for row in self.records),
                "partial": sum(row["status"] == "PARTIAL" for row in self.records),
                "unsupported": sum(row["status"] == "UNSUPPORTED" for row in self.records),
                "failed": sum(row["status"] == "FAILED" for row in self.records),
                "bytes_processed": self.total_bytes,
            },
            "records": self.records,
        }
        report_path = self.output_dir / "intake-report.json"
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        return report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Extract readable, traceable content without executing source files.")
    parser.add_argument("source", help="A file or directory to intake.")
    parser.add_argument("--output", default="intake-output", help="Output directory (default: intake-output).")
    args = parser.parse_args(argv)
    try:
        report = FileIntake(args.output).process(args.source)
    except (OSError, IntakeError) as exc:
        print(f"Intake failed: {exc}", file=sys.stderr)
        return 2
    print(json.dumps(report["summary"], ensure_ascii=False, indent=2))
    print(f"Report: {Path(args.output).resolve() / 'intake-report.json'}")
    return 0 if report["summary"]["failed"] == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
